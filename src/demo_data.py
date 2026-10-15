"""Sentetik (uydurma) zaman etüdü verisi üretir.

Amaç: Gerçek fabrika verisi paylaşılamadığında bile depoyu herkesin çalıştırabilmesi.
Üretilen Excel, gerçek formların yerleşimini (loader.py'nin beklediği) birebir taklit eder:
her iş öğesi satırının altında ham okumalar, sağda AK=ortalama / AL=%10 pay / AM=standart süre.

Tüm istasyon adları, süreler ve ürün kodu UYDURMADIR; gerçek bir işletmeyi temsil etmez.
Veri, kütüphanenin özelliklerini sergilemek için bilinçli olarak şu durumları içerir:
parti ölçümleri (5-35 adet), çok operatörlü istasyon, tek okumalı öğeler, rota (ERP) sayfası.

Kullanım:  python -m src.demo_data  [--output data/sample/demo_zaman_etudu.xlsx] [--seed 42]
"""
from __future__ import annotations

import argparse
import math
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook

PRODUCT = "DEMO100"
ALMA, BIRAKMA, ISLEM, TEST = "Alma- Yerleştirme", "Bırakma", "İşlem", "Test"

# (sayfa adı, kol soneki, parti, operatör, [(tür, öğe adı, parti başına ortalama sn, okuma sayısı, değişkenlik CV)])
STATIONS = [
    ("SIZINTI TESTİ", "", 1, 1, [(ALMA, "Aparata yerleştirme", 6.0, 3, .12), (TEST, "Sızıntı testi", 47.0, 2, .05),
                                  (ISLEM, "Prob kontrol", 10.5, 3, .20), (BIRAKMA, "Bırakma", 3.2, 3, .15)]),
    ("BİRLEŞTİRME KAYNAK", "", 1, 2, [(ISLEM, "Kaynak hazırlık", 25.5, 5, .35), (ISLEM, "Kaynak", 23.5, 5, .08),
                                       (ISLEM, "Gruplama", 9.2, 4, .10), (BIRAKMA, "Bırakma", 2.8, 4, .12)]),
    ("ETİKET BANTLAMA", "", 1, 1, [(ISLEM, "Bantlama", 19.8, 3, .10), (ISLEM, "Kalibrasyon", 7.1, 2, .08),
                                    (ISLEM, "Hat düzenleme", 6.8, 1, .0), (ALMA, "Parça alma", 2.5, 2, .10)]),
    ("ISITMA FIRINI", "", 5, 1, [(ISLEM, "Isıtma 1. adım", 80.0, 1, .0), (ISLEM, "Isıtma 2. adım", 58.0, 1, .0),
                                  (ALMA, "Parça alma", 12.0, 1, .0), (BIRAKMA, "Bırakma", 10.0, 1, .0)]),
    ("MONTAJ TAKMA", "", 1, 1, [(ALMA, "Parça alma", 4.5, 2, .10), (ISLEM, "Takma", 20.5, 3, .12),
                                 (ISLEM, "Sıkma", 4.1, 3, .10)]),
    ("-2 FORM VERME", "-2", 1, 1, [(ALMA, "Parça alma", 5.8, 3, .10), (ISLEM, "Form verme", 15.2, 3, .10),
                                    (BIRAKMA, "Bırakma", 4.1, 3, .12)]),
    ("PAKETLEME", "", 5, 1, [(ISLEM, "Klips takma", 52.0, 1, .0), (ALMA, "Parça alma", 30.0, 1, .0),
                              (BIRAKMA, "Kutuya bırakma", 29.0, 1, .0)]),
    ("ANA BORU KAYNAK", "", 6, 1, [(ALMA, "Parça alma", 38.0, 2, .10), (ISLEM, "Kaynak", 72.0, 2, .08),
                                    (BIRAKMA, "Bırakma", 12.0, 2, .10)]),
    ("-3 FORM VERME", "-3", 1, 1, [(ALMA, "Parça alma", 5.4, 2, .10), (ISLEM, "Form verme", 11.5, 2, .10),
                                    (BIRAKMA, "Bırakma", 2.0, 2, .15)]),
    ("-2 KESİM", "-2", 10, 1, [(ISLEM, "Kesim (10 adet)", 118.0, 2, .06), (ALMA, "Alma (10 adet)", 22.0, 2, .10)]),
    ("-3 SARIM", "-3", 1, 1, [(ISLEM, "Sarım", 11.0, 2, .10), (ALMA, "Parça alma", 2.2, 1, .0)]),
    ("-2 SARIM", "-2", 1, 1, [(ISLEM, "Sarım", 10.2, 2, .10), (ALMA, "Parça alma", 2.0, 1, .0)]),
    ("-4-5 HAVŞA", "-4-5", 10, 1, [(ISLEM, "Havşa (10 adet)", 100.0, 2, .08), (ALMA, "Alma (10 adet)", 22.0, 2, .10)]),
    ("KAPİLER TAKMA", "", 5, 1, [(ISLEM, "Takma", 40.0, 1, .0), (ALMA, "Parça alma", 8.0, 1, .0)]),
    ("-4 BÜKÜM", "-4", 1, 1, [(ISLEM, "Büküm", 6.0, 2, .10), (ALMA, "Parça alma", 1.8, 2, .10)]),
    ("-1 KAPİLER KAYNAK", "-1", 35, 1, [(ISLEM, "Kaynak (35 adet)", 255.0, 1, .0)]),
    ("BANTLAMA (2)", "", 1, 1, [(ISLEM, "Bantlama", 4.6, 3, .12), (ALMA, "Parça alma", 0.9, 3, .10)]),
    ("-1 BORU KESİM", "-1", 10, 1, [(ISLEM, "Kesim (10 adet)", 45.0, 2, .08), (ALMA, "Alma (10 adet)", 8.0, 2, .10)]),
    ("-3 KESİM MAKİNESİ", "-3", 7, 1, [(ISLEM, "Kesim (7 adet)", 30.0, 2, .08), (ALMA, "Alma (7 adet)", 5.0, 2, .10)]),
    ("SIVAMA", "", 10, 1, [(ISLEM, "Sıvama (10 adet)", 36.0, 2, .10), (ALMA, "Alma (10 adet)", 4.0, 2, .10)]),
    ("-4 BORU KESİM", "-4", 10, 1, [(ISLEM, "Kesim (10 adet)", 33.0, 2, .08), (ALMA, "Alma (10 adet)", 4.0, 2, .10)]),
    ("-4 CNC HAVŞA", "-4", 10, 1, [(ISLEM, "Havşa (10 adet)", 27.0, 2, .08), (ALMA, "Alma (10 adet)", 4.0, 2, .10)]),
    ("-6 HAVŞA", "-6", 10, 1, [(ISLEM, "Havşa (10 adet)", 22.0, 2, .08), (ALMA, "Alma (10 adet)", 4.0, 2, .10)]),
    ("BORU KESİM", "", 10, 1, [(ISLEM, "Kesim (10 adet)", 21.0, 2, .08), (ALMA, "Alma (10 adet)", 4.0, 2, .10)]),
    ("PAH KIRMA", "", 10, 1, [(ISLEM, "Pah kırma (10 adet)", 7.0, 2, .10), (ALMA, "Alma (10 adet)", 1.0, 2, .10)]),
]


def _write_station(wb: Workbook, spec, rng: np.random.Generator, date: datetime) -> None:
    name, suffix, batch, crew, elements = spec
    ws = wb.create_sheet(name)
    ws["A1"] = "ZAMAN ETÜDÜ FORMU (SENTETİK DEMO VERİSİ)"
    ws["A3"], ws["T3"], ws["AA3"], ws["AK3"] = "KOD", "OPERASYON", "NOT", "ANALİZ TARİHİ"
    ws["A4"] = PRODUCT + suffix
    ws["T4"] = name
    if crew > 1:
        ws["AA4"] = f"{crew} OPERATÖR ÇALIŞIYOR"
    ws["AK4"] = date
    ws["B5"], ws["C5"], ws["AK5"], ws["AL5"], ws["AM5"] = "TÜR", "İŞ ÖĞESİ", "ORTALAMA", "YORGUNLUK %10", "STANDART"

    row, total = 6, 0.0
    for op_type, element, base, n_read, cv in elements:
        if n_read == 1:
            readings = [round(base)]
        else:
            readings = [max(1, round(base * rng.lognormal(0, cv))) for _ in range(n_read)]
        mean = float(np.mean(readings))
        # Değerler (formül değil) yazılır: openpyxl data_only=True okumasında önbellek olmadığı için
        ws.cell(row, 2, op_type)
        ws.cell(row, 3, element)
        ws.cell(row, 37, mean)
        ws.cell(row, 38, mean * 0.10)
        ws.cell(row, 39, mean * 1.10)
        total += mean * 1.10
        for j, v in enumerate(readings):
            ws.cell(row + 1, 4 + j, v)
        row += 2

    if batch > 1:
        ws.cell(row + 1, 1, f"{batch} ürünün süresi alındı")
    total_ceil = math.ceil(total - 1e-9)
    ws.cell(row + 2, 1, f"STANDART SÜRE (yukarı yuvarlanmış) = {total_ceil}")
    ws.cell(row + 3, 1, round(total_ceil / batch, 2))


def build_demo_workbook(path: str | Path, seed: int = 42) -> Path:
    """Deterministik (seed'e bağlı) demo Excel dosyasını yazar ve yolunu döndürür."""
    rng = np.random.default_rng(seed)
    wb = Workbook()
    wb.remove(wb.active)
    date = datetime(2026, 1, 15)
    for spec in STATIONS:
        _write_station(wb, spec, rng, date)
    route = wb.create_sheet("Sayfa1")                 # ERP rota tablosu: analizde atlanır
    route.append(["ROTA", "SÜRE (dk)"])
    route.append([PRODUCT, 12])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--output", default="data/sample/demo_zaman_etudu.xlsx")
    ap.add_argument("--seed", type=int, default=42)
    a = ap.parse_args()
    print("Yazıldı:", build_demo_workbook(a.output, a.seed))
