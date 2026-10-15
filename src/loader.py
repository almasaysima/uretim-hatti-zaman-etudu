"""Excel zaman etüdü formlarını okuyup temiz (tidy) DataFrame'lere dönüştürür.

Formlar tablo değil "form" yapısındadır (her iş öğesi satırının altında ölçüm okumaları
bulunur), bu yüzden `pd.read_excel` yerine openpyxl ile ayrıştırılır; sonrası Pandas'tır.
Operatör / analist isimleri KASITLI olarak okunmaz (kişisel veri - KVKK).
"""
from __future__ import annotations

import math
import re
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from . import config as cfg


def fold(text) -> str:
    """Türkçe karakterleri sadeleştirip büyük harfe çevirir (karşılaştırma anahtarı)."""
    if text is None:
        return ""
    s = str(text).replace("ı", "i").replace("İ", "I")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return " ".join(s.upper().split())


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and not (
        isinstance(v, float) and math.isnan(v))


def categorize(op_type: str, element: str) -> str:
    """İş öğesini metot etüdü açısından 3 sınıfa ayırır."""
    t, e = fold(op_type), fold(element)
    if any(w in t for w in ("ALMA", "BIRAKMA", "YERLESTIRME")) or e in ("ALMA", "BIRAKMA"):
        return "Elleçleme (alma/bırakma/yerleştirme)"
    if "TEST" in t:
        return "Test"
    return "İşlem"


def _branch(code: str, product_code: str | None = None) -> str:
    """Ürün kodundan hat kolunu çıkarır: '<ürün>-1,2' → '-1,2'; yalnız '<ürün>' → 'Ana'."""
    base = re.escape(product_code or cfg.PRODUCT_CODE) if (product_code or cfg.PRODUCT_CODE) else r"[A-Za-z0-9]{4,}"
    m = re.search(rf"{base}\s*(-[\d,\-]+)?", code or "")
    return (m.group(1) or "Ana").strip() if m else "Diğer"


def _parse_sheet(ws, sheet_name: str, product_code: str | None = None):
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), max_col=40, values_only=True))
    meta_idx = next((i + 1 for i, r in enumerate(rows[:15]) if fold(r[0]).startswith("KOD")), None)
    if meta_idx is None:                       # zaman etüdü formu değil (ör. rota sayfası)
        return None
    meta = rows[meta_idx]
    code = str(meta[0] or "").strip()
    crew_match = re.search(r"(\d+)\s*OPERAT", fold(meta[cfg.COL_CREW_NOTE - 1]))
    crew = int(crew_match.group(1)) if crew_match else 1
    date = meta[cfg.COL_DATE - 1]

    elements, last_row = [], meta_idx
    for i in range(meta_idx + 1, len(rows) - 1):
        r = rows[i]
        ak, am = r[cfg.COL_MEAN - 1], r[cfg.COL_STD - 1]
        if r[cfg.COL_C - 1] is not None and _is_num(ak) and _is_num(am):
            readings = [float(v) for v in rows[i + 1][cfg.COL_FIRST_READING - 1:cfg.COL_LAST_READING]
                        if _is_num(v)]
            elements.append(dict(
                seq=len(elements) + 1,
                element=" ".join(str(r[cfg.COL_C - 1]).split()),
                op_type=" ".join(str(r[cfg.COL_B - 1] or "").split()),
                readings=readings, form_mean=float(ak), form_std=float(am)))
            last_row = i + 1

    batch, printed_total, printed_unit = 1, np.nan, np.nan
    for r in rows[last_row + 1:]:
        a = r[0]
        if isinstance(a, str):
            f = fold(a)
            mb = re.match(r"(\d+)\s+(MALZEME|URUN)", f)
            if mb:
                batch = int(mb.group(1))
            ms = re.search(r"STANDART SURE.*=\s*([\d.,]+)", f)
            if ms:
                printed_total = float(ms.group(1).replace(",", "."))
        elif _is_num(a):
            printed_unit = float(a)

    station = dict(station=sheet_name, code=code, branch=_branch(code, product_code),
                   op_name=" ".join(str(meta[cfg.COL_OP_NAME - 1] or "").split()),
                   crew=crew, batch=batch, form_total_std=printed_total,
                   form_unit_std=printed_unit,
                   study_date=pd.to_datetime(date, errors="coerce") if date is not None else pd.NaT)
    for e in elements:
        e["station"] = sheet_name
    return station, elements


def load_time_studies(xlsx_path: str | Path, product_code: str | None = None):
    """Döndürür: (elements_df, stations_df, skipped_sheets)."""
    wb = load_workbook(xlsx_path, data_only=True)
    stations, elements, skipped = [], [], []
    for ws in wb.worksheets:
        parsed = _parse_sheet(ws, ws.title, product_code)
        if parsed is None:
            skipped.append(ws.title)
            continue
        stations.append(parsed[0])
        elements.extend(parsed[1])
    st = pd.DataFrame(stations)
    st["order"] = range(1, len(st) + 1)
    return pd.DataFrame(elements), st, skipped


def _required_n(readings) -> float:
    """Niebel gözlem sayısı formülü: N' = (k·√(nΣx² − (Σx)²) / Σx)²  (k=40 → %95 / ±%5)."""
    n = len(readings)
    if n < 2:
        return np.nan
    x = np.asarray(readings, float)
    sx, sxx = x.sum(), (x ** 2).sum()
    inner = max(n * sxx - sx ** 2, 0.0)
    return float((cfg.SAMPLE_CONFIDENCE_K * math.sqrt(inner) / sx) ** 2)


def clean_and_compute(el: pd.DataFrame, st: pd.DataFrame) -> pd.DataFrame:
    """Ortalama → %10 yorgunluk payı → standart süre (parti ve birim başına)."""
    el = el.merge(st[["station", "batch"]], on="station", how="left").copy()
    el["n_readings"] = el["readings"].apply(len)
    # Ortalamayı ham okumalardan yeniden hesapla; okuma yoksa formdaki değere düş
    el["mean_s"] = [float(np.mean(r)) if r else fm for r, fm in zip(el["readings"], el["form_mean"])]
    el["fatigue_s"] = el["mean_s"] * cfg.FATIGUE_RATE
    el["std_s"] = el["mean_s"] + el["fatigue_s"]               # parti süresi (formdaki AM)
    el["std_unit_s"] = el["std_s"] / el["batch"]               # BİRİM başına standart süre
    el["mean_unit_s"] = el["mean_s"] / el["batch"]
    el["sd_s"] = [float(np.std(r, ddof=1)) if len(r) > 1 else np.nan for r in el["readings"]]
    el["cv"] = el["sd_s"] / el["mean_s"]
    el["category"] = [categorize(a, b) for a, b in zip(el["op_type"], el["element"])]
    el["form_match"] = np.isclose(el["std_s"], el["form_std"], atol=1e-6)
    el["n_required"] = [_required_n(r) for r in el["readings"]]
    el["sample_ok"] = [(n >= cfg.MIN_READINGS) and (np.isnan(req) or n >= req)
                       for n, req in zip(el["n_readings"], el["n_required"])]
    return el
